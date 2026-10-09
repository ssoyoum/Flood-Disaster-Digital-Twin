"""FloodOps 최종발표 자료 빌드 — 편집 가능한 PPTX (BASEMENT 2차 발표자료와 같은 도구).

    python -X utf8 build.py
    powershell -File finalize.ps1      # Noto Sans KR 포함 저장 + PDF

글자·사각형·원·선·화살표는 PowerPoint 기본 개체(Noto Sans KR)로 만든다.
단면도·아이콘·hero 패널 같은 그림만 SVG 로 넣고, PowerPoint 구버전·타 프로그램용 PNG 를 함께 둔다
(a:blip 에 asvg:svgBlip 확장, Office 2016+ 표준 방식).

글자 위치: PowerPoint 는 여백 0 텍스트 상자의 첫 기준선을 위 끝에서 0.96×글자크기, 기본 줄 간격을
1.2×글자크기로 그린다(2026-10-07 측정). 줄 간격 배수 m 이면 첫 기준선은 0.96 + 1.2(m-1) 이다.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from slides import SLIDES  # noqa: E402
from svgkit import FAMILY, H, W  # noqa: E402

OUT = HERE / "out"
DECK = HERE / "FloodOps_최종발표_v3.pptx"
# SVG 를 읽지 못하는 프로그램(한쇼·구버전·일부 미리보기)용: 그림을 PNG 로만 넣는다.
DECK_COMPAT = HERE / "FloodOps_최종발표_v3_호환용(PNG).pptx"
CHROME = Path("C:/Program Files/Google/Chrome/Application/chrome.exe")
E = 12192000 / W  # px → EMU
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
_svg_count = [0]


def emu(px: float) -> Emu:
    return Emu(int(round(px * E)))


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def drop_style(shape):
    """python-pptx 가 붙이는 테마 스타일(그림자·테마 선)을 뗀다. 색은 직접 지정한다."""
    style = shape._element.find(f"{{{P_NS}}}style")
    if style is not None:
        shape._element.remove(style)


def set_alpha(fill_parent_xml, opacity: float):
    clr = fill_parent_xml.find(f".//{{{A_NS}}}srgbClr")
    if clr is not None and opacity < 1:
        etree.SubElement(clr, f"{{{A_NS}}}alpha", val=str(int(opacity * 100000)))


def style_fill_line(shape, it):
    if it.get("fill", "none") == "none":
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(it["fill"])
        set_alpha(shape._element.spPr.find(f"{{{A_NS}}}solidFill"), it.get("opacity", 1.0))
    if it.get("stroke", "none") == "none":
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(it["stroke"])
        shape.line.width = emu(it.get("sw", 1))
        if it.get("dash"):
            shape.line.dash_style = MSO_LINE.DASH


def add_text(slide, it):
    size, lh = it["size"], it["lh"]
    nlines = it.get("nlines", len(it["rows"]))
    pad = 0 if it["wrap"] else 10
    w = it["w"] + pad
    x = it["x"] - w if it["anchor"] == "end" else it["x"] - w / 2 if it["anchor"] == "middle" else it["x"]
    top = it["y"] - size * (lh - 0.24)
    h = nlines * lh * size + 0.3 * size
    box = slide.shapes.add_textbox(emu(x), emu(top), emu(w), emu(h))
    first = "".join(s for s, _, _ in it["rows"][0])
    box.name = f"글 {first[:24]}"
    tf = box.text_frame
    # 자동 줄바꿈을 끈다. 여는 프로그램이 다른 글꼴로 대체해도 줄이 늘지 않게, 줄은 계산한 위치에서만 바꾼다.
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    align = {"start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT, "middle": PP_ALIGN.CENTER}[it["anchor"]]
    for i, row in enumerate(it["rows"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        # 한글을 글자 단위가 아니라 단어(공백) 단위로 줄바꿈한다.
        ppr = p._p.get_or_add_pPr()
        ppr.set("eaLnBrk", "0")
        ppr.set("latinLnBrk", "0")
        if abs(lh - 1.2) > 1e-6:
            p.line_spacing = lh / 1.2
        pieces = row
        if it.get("brk"):
            _, weight0, fill0 = row[0]
            pieces = []
            for j, line in enumerate(it["brk"]):
                pieces.append((line, weight0, fill0, j > 0))
        for piece in pieces:
            s, weight, fill = piece[:3]
            if len(piece) > 3 and piece[3]:
                p.add_line_break()
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = FAMILY
            f.size = Pt(size * 0.6)
            f.bold = weight >= 600
            f.color.rgb = rgb(fill)
            rpr = r._r.get_or_add_rPr()
            for tag in ("ea", "cs"):
                el = rpr.find(f"{{{A_NS}}}{tag}")
                if el is None:
                    el = etree.SubElement(rpr, f"{{{A_NS}}}{tag}")
                el.set("typeface", FAMILY)


def render_png(svg_path: Path, png_path: Path, w: float, h: float):
    sw, sh = max(int(round(w)), 1), max(int(round(h)), 1)
    subprocess.run([str(CHROME), "--headless=new", "--disable-gpu", "--hide-scrollbars", "--default-background-color=00000000",
                    f"--window-size={max(sw, 64)},{max(sh, 64)}", "--force-device-scale-factor=2",
                    f"--screenshot={png_path}", svg_path.resolve().as_uri()], check=True, capture_output=True)
    im = Image.open(png_path)
    im.crop((0, 0, sw * 2, sh * 2)).save(png_path)


def attach_svg(slide, pic, svg_bytes: bytes):
    _svg_count[0] += 1
    part = Part(PackURI(f"/ppt/media/art-{_svg_count[0]:03d}.svg"), "image/svg+xml", slide.part.package, svg_bytes)
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.find(f".//{{{A_NS}}}blip")
    ext_lst = etree.SubElement(blip, f"{{{A_NS}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{A_NS}}}ext", uri=SVG_EXT_URI)
    el = etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS})
    el.set(f"{{{R_NS}}}embed", rid)


def add_svg_picture(slide, svg_path: Path, png_path: Path, x, y, w, h, name, with_svg=True):
    pic = slide.shapes.add_picture(str(png_path), emu(x), emu(y), emu(w), emu(h))
    pic.name = name
    if with_svg:
        attach_svg(slide, pic, svg_path.read_bytes())


def prepare_shot(src: Path, w_px: float, h_px: float, dst: Path, crop=None) -> Path:
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
    target = w_px / h_px
    iw, ih = im.size
    if iw / ih > target + 1e-3:
        nw = int(ih * target)
        im = im.crop(((iw - nw) // 2, 0, (iw - nw) // 2 + nw, ih))
    elif iw / ih < target - 1e-3:
        im = im.crop((0, 0, iw, int(iw / target)))
    if im.width > 2400:
        im = im.resize((2400, int(2400 / target)), Image.LANCZOS)
    im.save(dst, "PNG", optimize=True)
    return dst


def main(deck: Path, with_svg: bool, only: set[int] | None = None):
    import slides
    slides.SOURCES.clear()
    _svg_count[0] = 0

    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(W), emu(H)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = "FloodOps — 과거 홍수를 재생하고 대응 시점을 비교하는 디지털 트윈"
    prs.core_properties.subject = "FloodOps 최종발표"
    prs.core_properties.language = "ko-KR"

    for i, fn in enumerate(SLIDES, start=1):
        if only and i not in only:
            continue
        canvas, images, notes = fn(i)
        name = f"{i:02d}-{fn.__name__.split('_', 1)[1]}"
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(canvas.bg)
        n_art = 0
        for it in canvas.items:
            k = it["kind"]
            if k in ("rect", "oval"):
                shape_type = MSO_SHAPE.OVAL if k == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if it.get("r") else MSO_SHAPE.RECTANGLE)
                sh = slide.shapes.add_shape(shape_type, emu(it["x"]), emu(it["y"]), emu(it["w"]), emu(it["h"]))
                drop_style(sh)
                if k == "rect" and it.get("r"):
                    sh.adjustments[0] = min(0.5, it["r"] / max(min(it["w"], it["h"]), 1))
                style_fill_line(sh, it)
                if it.get("name"):
                    sh.name = it["name"]
            elif k == "line":
                cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(it["x1"]), emu(it["y1"]), emu(it["x2"]), emu(it["y2"]))
                drop_style(cn)
                cn.line.color.rgb = rgb(it["stroke"])
                cn.line.width = emu(it["sw"])
                if it.get("dash"):
                    cn.line.dash_style = MSO_LINE.DASH
                if it.get("arrow"):
                    ln = cn.line._get_or_add_ln()
                    etree.SubElement(ln, f"{{{A_NS}}}tailEnd", type="triangle", w="med", len="med")
            elif k == "text":
                add_text(slide, it)
            elif k == "art":
                n_art += 1
                svg_path = OUT / "art" / f"{name}-{n_art}.svg"
                png_path = svg_path.with_suffix(".png")
                svg_path.write_text(it["art"].svg(), encoding="utf-8")
                if not png_path.exists():
                    render_png(svg_path, png_path, it["w"], it["h"])
                add_svg_picture(slide, svg_path, png_path, it["x"], it["y"], it["w"], it["h"], it["name"], with_svg)
            elif k == "svgfile":
                n_art += 1
                png_path = it["png"]
                if png_path is None or not png_path.exists():
                    png_path = OUT / "art" / f"{name}-{n_art}.png"
                    render_png(it["path"], png_path, it["w"], it["h"])
                add_svg_picture(slide, it["path"], png_path, it["x"], it["y"], it["w"], it["h"], it["name"], with_svg)
        for j, (src, x, y, w, h, crop) in enumerate(images):
            shot = prepare_shot((HERE / src).resolve(), w, h, OUT / "img" / f"{name}-{j}.png", crop)
            sp = slide.shapes.add_picture(str(shot), emu(x), emu(y), emu(w), emu(h))
            sp.name = f"화면 캡처 {j + 1}"
        slide.notes_slide.notes_text_frame.text = notes
        print(f"{name}: {len(canvas.items)} items, {n_art} art, {len(images)} image(s)")

    prs.save(deck)
    print(deck)


if __name__ == "__main__":
    for d in ("art", "img"):
        shutil.rmtree(OUT / d, ignore_errors=True)
        (OUT / d).mkdir(parents=True, exist_ok=True)
    for old in ("svg", "svg-text", "png"):
        shutil.rmtree(OUT / old, ignore_errors=True)
    # python build.py --only 12,14  → 해당 장만 담은 새 파일(기존 발표 파일은 건드리지 않음)
    if len(sys.argv) > 2 and sys.argv[1] == "--only":
        nums = sorted(int(n) for n in sys.argv[2].split(","))
        out = HERE / f"FloodOps_최종발표_v3_{'·'.join(str(n) for n in nums)}장_수정.pptx"
        main(out, with_svg=False, only=set(nums))
    else:
        main(DECK, with_svg=True)
        main(DECK_COMPAT, with_svg=False)
